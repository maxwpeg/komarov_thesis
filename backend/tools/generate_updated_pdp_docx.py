from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "docs"
OUTPUT_DIR = SOURCE_DIR / "ПДП_актуализация_2026-05-14"

DOCUMENTS = {
    "pmi": {
        "source": SOURCE_DIR / "PMI_komarov.docx",
        "output": OUTPUT_DIR / "ПМИ_комаров_ПДП_актуализировано.docx",
        "target": "ЛИСТ РЕГИСТРАЦИИ ИЗМЕНЕНИЙ",
    },
    "tp": {
        "source": SOURCE_DIR / "TP_komarov.docx",
        "output": OUTPUT_DIR / "ТП_комаров_ПДП_актуализировано.docx",
        "target": "3 СПИСОК ИСТОЧНИКОВ",
    },
    "ro": {
        "source": SOURCE_DIR / "RO_komarov_with_screenshots.docx",
        "output": OUTPUT_DIR / "РО_комаров_ПДП_актуализировано.docx",
        "target": "4 СООБЩЕНИЯ ОПЕРАТОРУ",
    },
}


COMMON_REPLACEMENTS = {
    "floorplan-ui": "frontend",
    "ReportLab==1.7.2": "reportlab==4.2.5",
    "pyfpdf / reportlab": "reportlab",
    "реляционная СУБД, рекомендуется PostgreSQL;": (
        "SQLite используется по умолчанию; PostgreSQL-совместимое подключение "
        "задается через переменную окружения DATABASE_URL;"
    ),
    "каталога `floorplan-ui`": "каталога `frontend`",
    "каталог `floorplan-ui`": "каталог `frontend`",
    "каталоге `floorplan-ui`": "каталоге `frontend`",
    "каталога floorplan-ui": "каталога frontend",
    "каталог floorplan-ui": "каталог frontend",
    "каталоге floorplan-ui": "каталоге frontend",
    "npm install`": "npm install`",
}


TP_REPLACEMENTS = {
    "На верхнем уровне репозитория выделены каталоги backend, frontend, floorplan, tests, docs, uploads, outputs и debug_output, а также набор Python-файлов, используемых для генерации проектной документации.": (
        "На верхнем уровне репозитория выделены две запускаемые части: каталог backend для серверной логики "
        "и каталог frontend для клиентского приложения. Внутри backend дополнительно находятся пакет backend, "
        "модуль floorplan, миграции alembic, тесты, инструменты и корневые Python-модули PDF-генерации."
    ),
    "• frontend – клиентское приложение React, реализующее интерфейс проектировщика.": (
        "• frontend – клиентское приложение React с package.json в корне каталога."
    ),
    "• tests – набор backend-интеграционных и API-тестов.": (
        "• backend/tests – набор backend-интеграционных и API-тестов."
    ),
    "• Файлы Project.py, PDFGenerator.py и связанные с ними страницы формируют контур PDF-генерации.": (
        "• Файлы backend/Project.py, backend/PDFGenerator.py и связанные с ними страницы формируют контур PDF-генерации."
    ),
    "Серверная часть запускается через файл main.py, который поднимает приложение backend.app:app.": (
        "Серверная часть запускается из каталога backend командой uvicorn backend.app:app --host 127.0.0.1 --port 8000."
    ),
    "Клиентская часть расположена в каталоге frontend и реализована на React 19. Файл frontend/src/App.jsx задает маршруты приложения и связывает страницы в единый пользовательский интерфейс.": (
        "Клиентская часть расположена в каталоге frontend и реализована на React 19. Файл frontend/src/App.jsx задает маршруты приложения, включая проекты, карточку проекта, редактор, каталог оборудования, дообучение, задачи, аудит и управление пользователями."
    ),
    "Доступ к серверному API централизован в файле frontend/src/api/client.js. В этом модуле выделены объекты projectsApi, equipmentApi, floorPlansApi, pipelineApi, recognitionApi, recognitionTrainingApi и elementsApi.": (
        "Доступ к серверному API централизован в файле frontend/src/api/client.js. В этом модуле выделены объекты projectsApi, equipmentApi, floorPlansApi, pipelineApi, recognitionApi, recognitionTrainingApi, backgroundTasksApi, auditApi и elementsApi."
    ),
    "Для запуска серверной части используется команда запуска Python-приложения через main.py, после чего FastAPI поднимает HTTP-сервис. Клиентская часть запускается из каталога frontend стандартными средствами React Scripts.": (
        "Для запуска серверной части пользователь переходит в каталог backend и выполняет uvicorn backend.app:app --host 127.0.0.1 --port 8000. Фоновый worker запускается из этого же каталога командой python tools/run_background_worker.py. Клиентская часть запускается из каталога frontend стандартными средствами React Scripts."
    ),
}


RO_REPLACEMENTS = {
    "В рекомендованном сценарии backend запускается локально на порту 8000, а frontend – на порту 3000. Для работы OCR и PDF-генерации в корне проекта должны сохраняться файлы шрифтов GOST_A.TTF и GOST_A_Bold.ttf, уже входящие в состав репозитория.": (
        "В рекомендованном сценарии backend запускается локально на порту 8000, а frontend – на порту 3000. "
        "Для работы OCR и PDF-генерации файлы шрифтов GOST_A.TTF и GOST_A_Bold.ttf, а также модели распознавания "
        "должны находиться в каталоге backend."
    ),
    "Перед первым запуском оператор или сопровождающий специалист должен развернуть зависимости backend и frontend. В корневой директории проекта выполняются следующие действия:": (
        "Перед первым запуском оператор или сопровождающий специалист должен развернуть зависимости backend и frontend. "
        "Команды выполняются из корней соответствующих каталогов:"
    ),
    "• установка серверных зависимостей командой `pip install -r backend/requirements.txt`;": (
        "• переход в каталог `backend` и установка серверных зависимостей командой `python -m pip install -r requirements.txt`;"
    ),
    "• переход в каталог `frontend` и установка клиентских зависимостей командой `npm install`;": (
        "• переход в каталог `frontend` и установка клиентских зависимостей командой `npm install`;"
    ),
    "В файле `.env` задаются значения `BOOTSTRAP_DEVELOPER_USERNAME`, `BOOTSTRAP_DEVELOPER_FULL_NAME` и `BOOTSTRAP_DEVELOPER_PASSWORD`.": (
        "В файле `backend/.env` задаются значения `BOOTSTRAP_DEVELOPER_USERNAME`, "
        "`BOOTSTRAP_DEVELOPER_FULL_NAME` и `BOOTSTRAP_DEVELOPER_PASSWORD`."
    ),
    "Для запуска клиентской части оператор переходит в каталог `frontend` и выполняет команду `npm start`.": (
        "Для запуска клиентской части оператор переходит в каталог `frontend` и выполняет команду `npm start`."
    ),
    "Генерация итоговой документации выполняется на странице проекта кнопкой «Сгенерировать PDF». Во время формирования интерфейс показывает состояние процесса, а по завершении инициирует скачивание готового файла.": (
        "Генерация итоговой документации выполняется на странице проекта кнопкой «Сгенерировать PDF». "
        "Во время формирования интерфейс показывает статус, стадию и процент выполнения фоновой задачи, "
        "а по завершении предоставляет ссылку на готовый файл."
    ),
}


PMI_REPLACEMENTS = {
    "Администратор: доступ ко всем проектам; управление пользователями и ролями; просмотр журналов аудита; служебные операции (перегенерация PDF, диагностика, управление хранилищем).": (
        "Администратор (роль developer): доступ ко всем проектам; управление пользователями и ролями; "
        "просмотр журналов аудита; служебные операции, включая повтор фоновых задач, перегенерацию PDF, "
        "диагностику и управление хранилищем."
    ),
    "Кнопка «Сгенерировать PDF» отправляет запрос на сервер, где происходит генерация PDF по определенным правилам (описано в тексте ВКР), ответ возвращается в виде готового файла, который автоматически загружается на устройство.": (
        "Кнопка «Сгенерировать PDF» отправляет запрос на сервер, где создается фоновая задача генерации PDF. "
        "Интерфейс показывает прогресс выполнения, а после завершения отображает ссылку на новую версию PDF-документа проекта."
    ),
}


def set_font(run, *, size: float = 14, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def normalize(value: str) -> str:
    return " ".join(value.split())


def replace_in_paragraph(paragraph: Paragraph, replacements: dict[str, str]) -> None:
    for run in paragraph.runs:
        text = run.text
        for old, new in replacements.items():
            text = text.replace(old, new)
        run.text = text
    full_text = paragraph.text
    updated = full_text
    for old, new in replacements.items():
        updated = updated.replace(old, new)
    if updated != full_text:
        paragraph.clear()
        run = paragraph.add_run(updated)
        is_heading = paragraph.style is not None and paragraph.style.name.startswith("Heading")
        set_font(run, bold=is_heading)


def replace_text(doc: Document, replacements: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph, replacements)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph, replacements)


def insert_before(paragraph: Paragraph, text: str = "", *, style: str | None = None, bold: bool = False) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style and style in paragraph.part.document.styles:
        new_para.style = paragraph.part.document.styles[style]
    if style and style.startswith("Heading"):
        new_para.paragraph_format.first_line_indent = None
        new_para.paragraph_format.space_before = Pt(6)
        new_para.paragraph_format.space_after = Pt(6)
    else:
        new_para.paragraph_format.first_line_indent = Cm(1.25)
        new_para.paragraph_format.space_before = Pt(0)
        new_para.paragraph_format.space_after = Pt(0)
        new_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    new_para.paragraph_format.line_spacing = 1.5
    run = new_para.add_run(text)
    set_font(run, bold=bold or bool(style and style.startswith("Heading")))
    return new_para


def insert_after(paragraph: Paragraph, text: str = "", *, style: str | None = None, bold: bool = False) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style and style in paragraph.part.document.styles:
        new_para.style = paragraph.part.document.styles[style]
    new_para.paragraph_format.first_line_indent = Cm(1.25)
    new_para.paragraph_format.line_spacing = 1.5
    new_para.paragraph_format.space_before = Pt(0)
    new_para.paragraph_format.space_after = Pt(0)
    new_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = new_para.add_run(text)
    set_font(run, bold=bold)
    return new_para


def find_target(doc: Document, marker: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if normalize(paragraph.text).startswith(marker):
            return paragraph
    raise RuntimeError(f"Target paragraph not found: {marker}")


def insert_block(doc: Document, marker: str, blocks: list[dict[str, object]]) -> None:
    target = find_target(doc, marker)
    for block in blocks:
        insert_before(
            target,
            str(block.get("text", "")),
            style=block.get("style") if isinstance(block.get("style"), str) else None,
            bold=bool(block.get("bold", False)),
        )


def normalize_sheet_counts(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = normalize(paragraph.text)
        if text.startswith("Листов:") or text.startswith("Листов "):
            paragraph.clear()
            run = paragraph.add_run("Листов ____")
            set_font(run)


def ensure_toc_placeholder(doc: Document) -> None:
    for index, paragraph in enumerate(doc.paragraphs):
        if normalize(paragraph.text) != "СОДЕРЖАНИЕ":
            continue
        following = next(
            (normalize(item.text) for item in doc.paragraphs[index + 1 : index + 4] if normalize(item.text)),
            "",
        )
        if "Оглавление обновляется" not in following:
            insert_after(paragraph, "Оглавление обновляется автоматически в текстовом редакторе.", style="Body Text")
        return


def heading(text: str) -> dict[str, object]:
    return {"text": text, "style": "Heading 1"}


def subheading(text: str) -> dict[str, object]:
    return {"text": text, "style": "Heading 2"}


def para(text: str) -> dict[str, object]:
    return {"text": text, "style": "Body Text"}


def bold_para(text: str) -> dict[str, object]:
    return {"text": text, "style": "Body Text", "bold": True}


def pmi_additions() -> list[dict[str, object]]:
    return [
        heading("Дополнительные методы испытаний актуализированной версии"),
        para(
            "Настоящий подраздел дополняет программу и методику испытаний с учетом фактической структуры проекта "
            "после разделения на запускаемые каталоги backend и frontend, а также с учетом реализованных "
            "административных функций, фоновой очереди, серверной проверки загрузок и версионирования PDF."
        ),
        bold_para("6.3.16 Проверка административного аудита и экспорта журнала"),
        para(
            "Назначение испытания. Испытание подтверждает, что пользователь с ролью developer может просматривать "
            "журнал аудита, применять фильтры по пользователю, сущности, действию и периоду, а также выгружать журнал "
            "в форматах CSV и JSON."
        ),
        para(
            "Порядок выполнения. Выполнить вход под учетной записью developer, открыть раздел «Аудит», задать фильтр "
            "по действию auth_login_succeeded или project_permanently_deleted, выполнить экспорт CSV, затем повторить "
            "экспорт в формате JSON через интерфейс или endpoint /api/audit/export."
        ),
        para(
            "Ожидаемый результат. В таблице отображаются только события, соответствующие выбранным условиям; файлы "
            "экспорта содержат идентификатор события, дату, пользователя, категорию, действие, связанную сущность и "
            "payload без раскрытия паролей и служебных cookie."
        ),
        bold_para("6.3.17 Проверка фоновых задач, прогресса и повторного запуска"),
        para(
            "Назначение испытания. Испытание подтверждает, что длительные операции запускаются асинхронно, возвращают "
            "HTTP 202 Accepted, отображают progress_percent и progress_stage, не блокируют интерфейс и могут быть "
            "повторены администратором при статусе failed."
        ),
        para(
            "Порядок выполнения. Запустить распознавание, генерацию PDF или дообучение, открыть страницу служебных "
            "задач, дождаться обновления статуса с периодичностью не реже одного раза в две секунды, после искусственно "
            "зафиксированной failed-задачи выполнить действие «Повторить»."
        ),
        para(
            "Ожидаемый результат. Задача проходит состояния queued, running и succeeded либо failed; прогресс изменяется "
            "от 0 до 100 процентов; повторный запуск failed-задачи возвращает ее в очередь и очищает сообщение об ошибке."
        ),
        bold_para("6.3.18 Проверка серверной валидации загружаемых планов"),
        para(
            "Назначение испытания. Испытание подтверждает, что ограничения на тип, размер и минимальное разрешение "
            "файла применяются не только в интерфейсе, но и на backend."
        ),
        para(
            "Порядок выполнения. Загрузить корректное изображение PNG/JPEG/JFIF/WebP размером не менее 200×200 пикселей; "
            "затем попытаться загрузить текстовый файл, поврежденное изображение и изображение меньшего разрешения."
        ),
        para(
            "Ожидаемый результат. Корректный план сохраняется и регистрируется в хранилище; неподдерживаемый тип "
            "отклоняется с кодом unsupported_upload_extension, поврежденный или слишком маленький файл — с кодом "
            "invalid_upload_file."
        ),
        bold_para("6.3.19 Проверка фильтров проектов и сохранения модели удаления"),
        para(
            "Назначение испытания. Испытание подтверждает, что список проектов фильтруется по автору, дате, статусу, "
            "объекту и номеру, при этом существующая модель soft-delete, корзины и permanent-delete сохраняется."
        ),
        para(
            "Порядок выполнения. Создать несколько проектов с различными объектами и владельцами, применить фильтры "
            "списка, переместить один проект в корзину, открыть список удаленных проектов и выполнить окончательное "
            "удаление под ролью developer."
        ),
        para(
            "Ожидаемый результат. В активном списке отображаются только проекты, удовлетворяющие фильтрам; удаленный "
            "проект доступен в корзине до permanent-delete; инженер не получает доступ к чужим проектам."
        ),
        bold_para("6.3.20 Проверка версий PDF-документов проекта"),
        para(
            "Назначение испытания. Испытание подтверждает, что каждая генерация PDF сохраняется как отдельная версия "
            "проекта, а не только как поле latest_pdf_path."
        ),
        para(
            "Порядок выполнения. Для проекта с подготовленными данными дважды запустить генерацию PDF, открыть карточку "
            "проекта и проверить список версий. Для каждой версии проверить дату генерации, автора, параметры, связь с "
            "фоновой задачей и ссылку скачивания."
        ),
        para(
            "Ожидаемый результат. В проекте отображается список PDF-версий; latest_pdf_path указывает на последнюю "
            "версию; ссылки на артефакты доступны только пользователям, имеющим доступ к проекту."
        ),
    ]


def tp_additions() -> list[dict[str, object]]:
    return [
        heading("2.11 Актуализация структуры и служебных подсистем"),
        para(
            "После актуализации проекта исходный код сгруппирован в две запускаемые папки. Серверная часть находится "
            "в каталоге backend: внутри него расположены пакет backend, модуль floorplan, тесты, миграции alembic, "
            "служебные инструменты и корневые модули PDF-генерации. Клиентская часть находится в каталоге frontend, "
            "где непосредственно расположен package.json."
        ),
        para(
            "Рабочие данные не относятся к тексту программы и хранятся отдельно: floor_plans.db, uploads, outputs, "
            "debug_output, storage_objects и локальные .env-файлы не переносятся в исходный код. Для использования "
            "старой рабочей базы после разделения каталогов применяется переменная DATA_DIR."
        ),
        para(
            "В серверной части добавлены маршруты /api/audit и /api/audit/export, которые реализуют административный "
            "журнал событий и выгрузку CSV/JSON. Модель AuditEvent хранит событие, категорию, пользователя, связанную "
            "сущность, request_id и payload. Доступ к журналу предоставляется только роли developer."
        ),
        para(
            "Фоновые операции вынесены в persistent-очередь BackgroundTask. Для задач распознавания, генерации PDF "
            "и дообучения сохраняются status, attempts, progress_percent, progress_stage, heartbeat_at и result_payload. "
            "Сервис BackgroundTaskService поддерживает постановку в очередь, claim, heartbeat, progress update, fail, "
            "cancel и retry failed-задач."
        ),
        para(
            "PDF-документы проекта представлены отдельной сущностью ProjectPdfVersion. Она хранит путь к файлу, автора "
            "генерации, параметры, дату создания и связь с фоновой задачей. Благодаря этому проект сохраняет историю "
            "версий PDF, а не только ссылку на последний сформированный файл."
        ),
        para(
            "Загрузка планов этажа ограничена изображениями PNG, JPG, JPEG, JFIF и WebP. Минимальное разрешение плана "
            "проверяется на backend и составляет 200×200 пикселей. PDF-файлы используются как документы и артефакты, "
            "но не импортируются как изображение плана этажа в текущей версии программы."
        ),
        para(
            "Пароли пользователей хэшируются по схеме PBKDF2-SHA256. Генерация проектных PDF выполняется библиотекой "
            "reportlab версии 4.2.5. Основная конфигурация базы данных использует SQLite по умолчанию и сохраняет "
            "совместимость с PostgreSQL через переменную DATABASE_URL."
        ),
        para(
            "Проверка текста программы выполняется командами: из каталога backend — python -m compileall backend "
            "floorplan tools и python -m pytest -q; из каталога frontend — npm.cmd test -- --watch=false --runInBand "
            "и npm.cmd run build."
        ),
    ]


def ro_execution_additions() -> list[dict[str, object]]:
    return [
        heading("3.11 Работа с аудитом, фоновыми задачами и версиями PDF"),
        para(
            "Пользователь с ролью developer имеет доступ к разделу «Аудит». В этом разделе отображается журнал "
            "значимых действий: входы в систему, изменения пользователей, операции с проектами, оборудованием, "
            "запуск генерации PDF, дообучение и другие события. Оператор-разработчик может отфильтровать записи по "
            "пользователю, проекту, категории, действию и периоду."
        ),
        para(
            "Для передачи журнала в отчет или для последующего анализа используются кнопки экспорта CSV и JSON. "
            "Файл экспорта содержит технические поля события, но не содержит паролей и не раскрывает значение "
            "серверной cookie-сессии."
        ),
        para(
            "Раздел фоновых задач показывает длительные операции, которые выполняются worker-процессом. Для каждой "
            "задачи отображаются тип, статус, инициатор, связанный проект или план, процент выполнения и текущая "
            "стадия. Интерфейс обновляет данные с интервалом около двух секунд."
        ),
        para(
            "Если задача завершилась ошибкой, пользователь с ролью developer может выполнить повторный запуск. "
            "Повтор доступен только для failed-задач; отмена доступна только для задач в состоянии queued. Такое "
            "поведение предотвращает повреждение уже выполняемых операций."
        ),
        para(
            "В карточке проекта после генерации документации отображается список PDF-версий. Для версии указываются "
            "дата формирования, автор, параметры генерации, идентификатор фоновой задачи и ссылка скачивания. "
            "Последняя версия одновременно используется как актуальный PDF проекта."
        ),
        para(
            "При загрузке плана этажа оператор видит имя файла, тип, размер и состояние загрузки. Если файл имеет "
            "неподдерживаемый тип, поврежден или меньше 200×200 пикселей, backend отклоняет загрузку, а интерфейс "
            "показывает диагностическое сообщение."
        ),
    ]


def ro_message_additions() -> list[dict[str, object]]:
    return [
        heading("4.12 invalid_upload_file, unsupported_upload_extension и ошибки фоновых задач"),
        para(
            "Сообщение `unsupported_upload_extension` означает, что оператор попытался загрузить файл с расширением, "
            "не разрешенным для планов этажа. В текущей версии для планов используются изображения PNG, JPG, JPEG, "
            "JFIF и WebP. PDF-планы не импортируются как планы этажа."
        ),
        para(
            "Сообщение `invalid_upload_file` означает, что файл не удалось распознать как корректное изображение "
            "или его разрешение меньше 200×200 пикселей. Оператору следует подготовить изображение достаточного "
            "качества и повторить загрузку."
        ),
        para(
            "Сообщение `background_task_not_retryable` возникает, если оператор пытается повторить задачу, которая "
            "не находится в статусе failed. Сообщение `background_task_not_cancelable` означает, что отмена невозможна, "
            "поскольку задача уже выполняется или завершена."
        ),
        para(
            "Если при скачивании файла появляется сообщение о недоступности артефакта, оператор должен проверить, "
            "что он работает с проектом, к которому имеет доступ, и что файл не был удален из runtime-хранилища."
        ),
    ]


def update_document(kind: str) -> Path:
    spec = DOCUMENTS[kind]
    doc = Document(str(spec["source"]))
    replacements = dict(COMMON_REPLACEMENTS)
    if kind == "tp":
        replacements.update(TP_REPLACEMENTS)
    elif kind == "ro":
        replacements.update(RO_REPLACEMENTS)
    elif kind == "pmi":
        replacements.update(PMI_REPLACEMENTS)
    replace_text(doc, replacements)
    normalize_sheet_counts(doc)
    ensure_toc_placeholder(doc)

    if kind == "pmi":
        additions = pmi_additions()
        insert_block(doc, str(spec["target"]), additions)
    elif kind == "tp":
        additions = tp_additions()
        insert_block(doc, str(spec["target"]), additions)
    else:
        insert_block(doc, str(spec["target"]), ro_execution_additions())
        insert_block(doc, "5 СПИСОК ИСТОЧНИКОВ", ro_message_additions())

    doc.core_properties.title = f"{doc.core_properties.title or ''} Актуализированная версия".strip()
    doc.core_properties.comments = (
        "Актуализировано без изменения исходных документов. "
        f"Дата формирования: {date.today().isoformat()}."
    )
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = spec["output"]
    doc.save(str(output))
    return output


def main() -> None:
    for kind in ("pmi", "tp", "ro"):
        output = update_document(kind)
        print(output)


if __name__ == "__main__":
    main()
