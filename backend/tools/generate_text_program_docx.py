from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "docs" / "TP_komarov.docx"

PROGRAM_NAME_RU = (
    "Веб-приложение для проектирования пожарной сигнализации и автоматического "
    "формирования проектной документации"
)
PROGRAM_NAME_EN = (
    "Web Application for Fire Alarm Design and Automatic Generation of Design Documentation"
)
PROGRAM_SHORT_NAME = "PlanVision"
DOCUMENT_CODE_APPROVAL = "RU.17701729.08.03-01 ТП 01-1-ЛУ"
DOCUMENT_CODE_MAIN = "RU.17701729.08.03-01 ТП 01-1"
CITY_AND_YEAR = "Москва 2026"
STUDENT_GROUP = "БПИ225"
STUDENT_NAME = "М. С. Комаров"


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("left", "top", "right", "bottom", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f"w:{edge}"
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key, value in edge_data.items():
                element.set(qn(f"w:{key}"), str(value))


def set_table_borders(table, *, visible: bool):
    spec = {"val": "single", "sz": "4", "color": "000000"} if visible else {"val": "nil"}
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top=spec,
                bottom=spec,
                left=spec,
                right=spec,
                insideH=spec,
                insideV=spec,
            )


def set_run_font(run, *, size: float = 14, bold: bool | None = None, italic: bool | None = None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def apply_body_style(paragraph):
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    align: WD_ALIGN_PARAGRAPH | None = None,
    bold: bool = False,
    italic: bool = False,
    size: float = 14,
    first_line_indent_cm: float | None = 1.25,
    line_spacing: float = 1.5,
    uppercase: bool = False,
):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    if first_line_indent_cm is None:
        paragraph.paragraph_format.first_line_indent = None
    else:
        paragraph.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    content = text.upper() if uppercase else text
    run = paragraph.add_run(content)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return paragraph


def add_heading(doc: Document, text: str, *, level: int = 1):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Heading 1"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=14, bold=True)
    return paragraph


def add_bullet(doc: Document, text: str):
    paragraph = add_paragraph(
        doc,
        f"• {text}",
        first_line_indent_cm=None,
        line_spacing=1.5,
    )
    paragraph.paragraph_format.left_indent = Cm(0.75)
    paragraph.paragraph_format.first_line_indent = Cm(-0.5)
    return paragraph


def add_page_break(doc: Document):
    doc.add_page_break()


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Оглавление обновляется автоматически в текстовом редакторе."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, size=14)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    for style_name in ("Normal", "Body Text", "Normal (Web)", "List Paragraph"):
        if style_name in doc.styles:
            font = doc.styles[style_name].font
            font.name = "Times New Roman"
            font.size = Pt(14)
    if "Heading 1" in doc.styles:
        font = doc.styles["Heading 1"].font
        font.name = "Times New Roman"
        font.size = Pt(14)
        font.bold = True

    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.5
    normal.space_after = Pt(0)
    normal.space_before = Pt(0)


def add_approval_page(doc: Document):
    add_paragraph(doc, CITY_AND_YEAR, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line_indent_cm=None)
    add_paragraph(
        doc,
        "ПРАВИТЕЛЬСТВО РОССИЙСКОЙ ФЕДЕРАЦИИ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        first_line_indent_cm=None,
    )
    add_paragraph(
        doc,
        "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        first_line_indent_cm=None,
    )
    add_paragraph(
        doc,
        "ВЫСШЕГО ОБРАЗОВАНИЯ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        first_line_indent_cm=None,
    )
    add_paragraph(
        doc,
        "«НАЦИОНАЛЬНЫЙ ИССЛЕДОВАТЕЛЬСКИЙ УНИВЕРСИТЕТ",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        first_line_indent_cm=None,
    )
    add_paragraph(
        doc,
        "«ВЫСШАЯ ШКОЛА ЭКОНОМИКИ»",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        first_line_indent_cm=None,
    )
    add_paragraph(doc, "", first_line_indent_cm=None)
    add_paragraph(doc, "Факультет компьютерных наук", align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent_cm=None)
    add_paragraph(
        doc,
        "Образовательная программа «Программная инженерия»",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent_cm=None,
    )

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, visible=False)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    left.text = ""
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("СОГЛАСОВАНО")
    set_run_font(r, size=14, bold=True)
    for line in (
        "Руководитель выпускной квалификационной работы",
        "___________________ / __________________ /",
        "«___» _____________ 2026 г.",
    ):
        q = left.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rr = q.add_run(line)
        set_run_font(rr, size=14)

    right.text = ""
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("УТВЕРЖДАЮ")
    set_run_font(r, size=14, bold=True)
    for line in (
        "Академический руководитель образовательной программы",
        "«Программная инженерия»",
        "___________________ / __________________ /",
        "«___» _____________ 2026 г.",
    ):
        q = right.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rr = q.add_run(line)
        set_run_font(rr, size=14)

    for _ in range(6):
        add_paragraph(doc, "", first_line_indent_cm=None)

    add_paragraph(
        doc,
        PROGRAM_NAME_RU,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        first_line_indent_cm=None,
        uppercase=True,
    )
    add_paragraph(doc, "Текст программы", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_line_indent_cm=None)
    add_paragraph(doc, "ЛИСТ УТВЕРЖДЕНИЯ", align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent_cm=None)
    add_paragraph(doc, DOCUMENT_CODE_APPROVAL, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_line_indent_cm=None)
    add_paragraph(doc, "", first_line_indent_cm=None)
    add_paragraph(doc, "Исполнитель", align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_cm=None)
    add_paragraph(
        doc,
        f"Студент группы {STUDENT_GROUP}                     ___________ / {STUDENT_NAME} /",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent_cm=None,
    )
    add_paragraph(doc, "«___» _____________ 2026 г.", align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent_cm=None)


def add_title_page(doc: Document):
    add_paragraph(doc, CITY_AND_YEAR, align=WD_ALIGN_PARAGRAPH.RIGHT, first_line_indent_cm=None)
    add_paragraph(doc, "", first_line_indent_cm=None)
    add_paragraph(
        doc,
        f"УТВЕРЖДЕН    {DOCUMENT_CODE_APPROVAL}",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        bold=True,
        first_line_indent_cm=None,
    )
    add_paragraph(doc, "", first_line_indent_cm=None)
    add_paragraph(
        doc,
        PROGRAM_NAME_RU,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        first_line_indent_cm=None,
        uppercase=True,
    )
    add_paragraph(doc, "Текст программы", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_line_indent_cm=None)
    add_paragraph(doc, DOCUMENT_CODE_MAIN, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, first_line_indent_cm=None)
    add_paragraph(doc, "Листов ____", align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent_cm=None)


def add_introduction(doc: Document):
    add_heading(doc, "1 ВВЕДЕНИЕ")
    add_heading(doc, "1.1 Наименование программы")
    add_paragraph(doc, f"Наименование программы – «{PROGRAM_NAME_RU}».")
    add_paragraph(doc, f"Наименование программы на английском языке – «{PROGRAM_NAME_EN}».")
    add_paragraph(doc, f"Краткое наименование программы – «{PROGRAM_SHORT_NAME}».")

    add_heading(doc, "1.2 Краткая характеристика области применения программы")
    add_paragraph(
        doc,
        "Программа предназначена для автоматизации задач, возникающих при подготовке "
        "проектной документации по системам пожарной сигнализации и связанным подсистемам "
        "на основе поэтажных планов зданий. Система используется как инженерный "
        "веб-инструмент, объединяющий хранение проектов, редактирование графических планов, "
        "распознавание архитектурных элементов, размещение оборудования, формирование "
        "чертежей и выпуск выходных документов.",
    )
    add_paragraph(
        doc,
        "Область применения программы охватывает учебные и прикладные сценарии "
        "проектирования: подготовку и ведение проектных карточек, загрузку планов этажей, "
        "полуавтоматическое построение цифровой модели здания, размещение пожарных "
        "извещателей, приборов и устройств СОУЭ, а также генерацию итоговой документации "
        "в формате PDF. Дополнительно программа поддерживает сбор исправленных пользователем "
        "примеров для последующего дообучения моделей распознавания.",
    )
    add_paragraph(
        doc,
        "По своему назначению программа является многокомпонентной системой. В ее составе "
        "совместно работают серверная часть на FastAPI, клиентское приложение на React, "
        "модуль компьютерного зрения для обработки изображений планов, подсистема "
        "геометрических преобразований и контур формирования проектной документации по "
        "требованиям ГОСТ.",
    )


def add_program_text(doc: Document):
    add_heading(doc, "2 ТЕКСТ ПРОГРАММЫ")
    add_paragraph(
        doc,
        "Полный текст программы хранится в репозитории разработки проекта "
        "«komarov_thesis» и представлен исходными файлами на языках Python и "
        "JavaScript/JSX. Исходный код организован по функциональным подсистемам, каждая "
        "из которых отвечает за отдельный этап работы пользователя: ведение проекта, "
        "редактирование плана, распознавание, проектирование оборудования, выпуск "
        "документации и сопровождение обучающих данных.",
    )

    add_heading(doc, "2.1 Общая организация репозитория")
    add_paragraph(
        doc,
        "На верхнем уровне репозитория выделены каталоги backend, frontend, floorplan, "
        "tests, docs, uploads, outputs и debug_output, а также набор Python-файлов, "
        "используемых для генерации проектной документации. Такая структура отражает "
        "разделение программы на активную веб-часть, вычислительные модули и артефакты "
        "работы системы.",
    )
    add_bullet(doc, "backend – серверное приложение FastAPI, роутеры, бизнес-логика, ORM и прикладные сервисы.")
    add_bullet(doc, "frontend – клиентское приложение React, реализующее интерфейс проектировщика.")
    add_bullet(doc, "floorplan – отдельный Python-модуль распознавания поэтажных планов и геометрической обработки.")
    add_bullet(doc, "tests – набор backend-интеграционных и API-тестов.")
    add_bullet(doc, "docs – вспомогательная документация по архитектуре, распознаванию и отчетным материалам.")
    add_bullet(doc, "uploads, outputs, debug_output – рабочие каталоги для входных файлов, результатов и отладочных артефактов.")
    add_bullet(doc, "Файлы Project.py, PDFGenerator.py и связанные с ними страницы формируют контур PDF-генерации.")

    add_heading(doc, "2.2 Серверная часть программы")
    add_paragraph(
        doc,
        "Серверная часть запускается через файл main.py, который поднимает приложение "
        "backend.app:app. В файле backend/app.py реализована фабрика FastAPI-приложения: "
        "создаются рабочие директории, инициализируется база данных, регистрируются шрифты "
        "для PDF-генерации, подключаются middleware, статические каталоги и OpenAPI-описание.",
    )
    add_paragraph(
        doc,
        "Маршрутизация запросов сосредоточена в каталоге backend/routers. На текущий момент "
        "в программе выделены маршруты auth, users, projects, equipment, floor_plans, "
        "elements, pipeline, recognition, recognition_training и pdf. Такое разбиение "
        "позволяет отделить пользовательские сценарии друг от друга и не концентрировать "
        "всю прикладную логику в одном контроллере.",
    )
    add_paragraph(
        doc,
        "Основная предметная логика в новой версии backend организована по модульному "
        "принципу в каталоге backend/modules. В нем выделены подсистемы projects, "
        "floor_plans, equipment, elements_geometry, pipeline, recognition, signal_design, "
        "documents и shared. Внутри модулей используется слоение application, domain и "
        "infrastructure, что облегчает сопровождение кода и локализует изменения.",
    )
    add_paragraph(
        doc,
        "Дополнительно в проекте сохраняется слой legacy-сервисов в каталоге "
        "backend/services. Он содержит код, который все еще реально используется системой, "
        "в частности при работе пайплайна, распознавания и дообучения. Переход к полностью "
        "модульной архитектуре выполняется постепенно, поэтому новая и старая модели "
        "организации кода сейчас сосуществуют в рамках одного серверного контура.",
    )
    add_paragraph(
        doc,
        "Отдельный блок серверной части образует контур аутентификации и управления "
        "пользователями. Файл backend/auth.py реализует хеширование паролей по схеме "
        "PBKDF2-SHA256, создание серверных сессий, хранение cookie и проверку ролей "
        "developer и engineer. Маршруты backend/routers/auth.py и "
        "backend/routers/users.py обеспечивают вход в систему, выход, получение профиля, "
        "создание и редактирование пользователей, а также контроль сохранения хотя бы "
        "одного активного разработчика.",
    )
    add_paragraph(
        doc,
        "Обмен данными между сервером и клиентом описывается Pydantic-схемами. "
        "SQLAlchemy-модели и репозитории используются для хранения проектов, этажей, "
        "элементов плана, состояния этапов pipeline, пользовательских сессий, обучающих "
        "примеров и других сущностей предметной области.",
    )

    add_heading(doc, "2.3 Клиентская часть программы")
    add_paragraph(
        doc,
        "Клиентская часть расположена в каталоге frontend и реализована на React 19. "
        "Файл frontend/src/App.jsx задает маршруты приложения и связывает страницы "
        "в единый пользовательский интерфейс. Через него подключаются страницы списка "
        "проектов, создания проекта, карточки проекта, редактора этажа, каталога "
        "оборудования, страницы дообучения, входа в систему и управления пользователями.",
    )
    add_paragraph(
        doc,
        "Доступ к серверному API централизован в файле frontend/src/api/client.js. "
        "В этом модуле выделены объекты projectsApi, equipmentApi, floorPlansApi, "
        "pipelineApi, recognitionApi, recognitionTrainingApi и elementsApi. Благодаря такой "
        "централизации запросы не дублируются по отдельным компонентам и сохраняется единый "
        "подход к обработке ответов и ошибок.",
    )
    add_paragraph(
        doc,
        "Ключевым пользовательским модулем является страница FloorPlanEditor.jsx. Она "
        "реализует визуализацию плана через React-Konva, масштабирование и панорамирование "
        "холста, выделение и редактирование графических элементов, работу с пожарными "
        "извещателями и кабельными трассами, а также оркестрацию шагов pipeline. По сути "
        "именно этот компонент является главным рабочим местом проектировщика.",
    )
    add_paragraph(
        doc,
        "Дополнительные страницы выполняют специализированные функции. ProjectList.jsx "
        "отвечает за просмотр перечня проектов, CreateProject.jsx – за создание новой "
        "карточки проекта, ProjectDetail.jsx – за работу с этажами, документами и "
        "генерацией PDF, EquipmentCatalogPage.jsx – за каталог оборудования, "
        "RecognitionTrainingPage.jsx – за управление обучающими примерами, LoginPage.jsx – "
        "за пользовательский вход, а UsersPage.jsx – за администрирование учетных записей.",
    )

    add_heading(doc, "2.4 Модуль распознавания поэтажных планов")
    add_paragraph(
        doc,
        "Отдельный каталог floorplan содержит алгоритмическое ядро обработки графических "
        "планов. Файл floorplan/main.py связывает этапы распознавания в последовательный "
        "сценарий. В состав модуля также входят preprocess.py, walls.py, walls_morph.py, "
        "openings.py, rooms.py, geometry.py, ocr_dimensions.py, visualize.py, "
        "floorplan_types.py и examples.py.",
    )
    add_paragraph(
        doc,
        "Назначение этих файлов распределено по этапам. preprocess.py подготавливает "
        "изображение к дальнейшей обработке; walls.py и walls_morph.py выявляют стены и "
        "структурные линии; openings.py определяет проемы; rooms.py сегментирует помещения "
        "и вычисляет их характеристики; ocr_dimensions.py извлекает размерные подписи; "
        "geometry.py содержит геометрические операции и вспомогательные функции; "
        "visualize.py формирует наглядные промежуточные результаты.",
    )
    add_paragraph(
        doc,
        "Результаты работы модуля распознавания не считаются окончательными без участия "
        "пользователя. После автоматического выделения элементов они передаются в редактор "
        "плана, где инженер может уточнить геометрию, исправить ошибки и сохранить "
        "скорректированное состояние. Такой подход позволяет совместить скорость "
        "автоматической обработки с точностью экспертной корректировки.",
    )

    add_heading(doc, "2.5 Модуль проектирования и генерации документации")
    add_paragraph(
        doc,
        "Формирование выходной проектной документации реализовано в корневом наборе "
        "Python-файлов, который исторически возник раньше модульного backend и до сих пор "
        "используется системой. Центральную роль играют Project.py и PDFGenerator.py. "
        "Они собирают комплект листов, создают PDF-файл и подготавливают его к выдаче "
        "пользователю.",
    )
    add_paragraph(
        doc,
        "Вспомогательные файлы реализуют отдельные страницы документа. pdf_documents/TitlePage.py "
        "формирует титульные листы, pdf_documents/DrawingPage.py – чертежи с наложением проектных "
        "элементов, pdf_documents/GeneralDataPage.py и pdf_documents/GeneralInstructionsPage.py – общие данные и "
        "указания, pdf_documents/ConventionalSymbolsPage.py – условные обозначения, "
        "pdf_documents/SpecificationPage.py – спецификацию оборудования, "
        "pdf_documents/PowerConsumptionCalculationPage.py – расчетные таблицы, "
        "pdf_documents/AdditionalInfoPage.py – дополнительную справочную информацию.",
    )
    add_paragraph(
        doc,
        "Для формирования страниц используются ReportLab, графические ресурсы и шрифты, "
        "включая GOST_A.TTF. Это позволяет генерировать документы, максимально близкие "
        "по внешнему виду к традиционным печатным листам проектной документации.",
    )

    add_heading(doc, "2.6 Подсистема хранения данных и рабочих артефактов")
    add_paragraph(
        doc,
        "Базой данных по умолчанию служит файл floor_plans.db. Для доступа к данным "
        "используется SQLAlchemy. В БД хранятся карточки проектов, сведения об этажах, "
        "геометрические элементы, pipeline-состояния, сведения об оборудовании, данные "
        "обучающих примеров, пользовательские учетные записи и серверные сессии.",
    )
    add_paragraph(
        doc,
        "Каталог uploads предназначен для загруженных изображений и других входных файлов. "
        "Каталог outputs хранит выходные документы, сформированные PDF-файлы, экспортированные "
        "датасеты и артефакты обучения. Каталог debug_output используется для промежуточных "
        "изображений и отладочных результатов распознавания. Такое разделение упрощает "
        "поиск файлов и снижает риск смешивания исходных данных с результатами обработки.",
    )
    add_paragraph(
        doc,
        "Для эволюции схемы хранения в проекте предусмотрен каталог alembic с миграциями. "
        "Это позволяет изменять структуру БД без ручного пересоздания таблиц и обеспечивает "
        "повторяемость развертывания программы на новой среде.",
    )

    add_heading(doc, "2.7 Подсистема дообучения распознавания")
    add_paragraph(
        doc,
        "В проект встроен самостоятельный контур дообучения распознавания на основе "
        "исправлений, сделанных пользователем. Когда инженер корректирует стены или "
        "проемы после прохождения соответствующего шага pipeline, система может сохранить "
        "разницу между исходным и исправленным состоянием как обучающий пример.",
    )
    add_paragraph(
        doc,
        "Сбор и обработка таких данных выполняется сервисами "
        "backend/services/recognition_training_feedback_service.py и "
        "backend/services/recognition_training_management_service.py. "
        "Публичный API предоставляется маршрутом backend/routers/recognition_training.py, "
        "а отдельный запуск обучения выполняется сценарием tools/run_recognition_training.py.",
    )
    add_paragraph(
        doc,
        "Логически подсистема оперирует тремя сущностями: обучающий пример, пакет "
        "обучающих данных и запуск обучения. Экспорт примеров размещается в каталоге "
        "outputs/recognition_feedback, а результаты конкретных запусков – в "
        "outputs/recognition_training. Подготовленный датасет имеет формат, пригодный для "
        "обучения через Ultralytics YOLO. При этом новая модель не подменяет рабочую "
        "автоматически, а сохраняется как кандидат для дальнейшей проверки.",
    )

    add_heading(doc, "2.8 Порядок запуска и основные входные точки")
    add_paragraph(
        doc,
        "Для запуска серверной части используется команда запуска Python-приложения через "
        "main.py, после чего FastAPI поднимает HTTP-сервис. Клиентская часть запускается "
        "из каталога frontend стандартными средствами React Scripts. При таком сценарии "
        "frontend работает как отдельное приложение и обращается к backend по API.",
    )
    add_paragraph(
        doc,
        "Тестирование серверной части выполняется через pytest, а клиентской – средствами "
        "React Testing Library и Jest-совместимого окружения, подключенного через npm-скрипты. "
        "Отдельные тесты проекта покрывают карточки проектов, редактор плана, каталог "
        "оборудования, дообучение распознавания и интеграцию API-клиента.",
    )
    add_paragraph(
        doc,
        "Таким образом, текст программы представляет собой не один монолитный исполняемый "
        "файл, а совокупность согласованно работающих модулей. Их совместная работа "
        "обеспечивает полный цикл от загрузки поэтажного плана и его распознавания до "
        "подготовки инженерного проекта и выпуска оформленной документации.",
    )


def add_sources(doc: Document):
    add_heading(doc, "3 СПИСОК ИСТОЧНИКОВ")
    sources = [
        "1) ГОСТ 19.101–77 Виды программ и программных документов. // Единая система программной документации.",
        "2) ГОСТ 19.103–77 Обозначения программ и программных документов. // Единая система программной документации.",
        "3) ГОСТ 19.104–78 Основные надписи. // Единая система программной документации.",
        "4) ГОСТ 19.105–78 Общие требования к программным документам. // Единая система программной документации.",
        "5) ГОСТ 19.106–78 Требования к программным документам, выполненным печатным способом. // Единая система программной документации.",
        "6) ГОСТ 19.401–78 Текст программы. Требования к содержанию и оформлению. // Единая система программной документации.",
        "7) Python. [Электронный ресурс]. URL: https://www.python.org/ (дата обращения: 22.04.2026).",
        "8) FastAPI. [Электронный ресурс]. URL: https://fastapi.tiangolo.com/ (дата обращения: 22.04.2026).",
        "9) React. [Электронный ресурс]. URL: https://react.dev/ (дата обращения: 22.04.2026).",
        "10) SQLAlchemy. [Электронный ресурс]. URL: https://www.sqlalchemy.org/ (дата обращения: 22.04.2026).",
        "11) OpenCV. [Электронный ресурс]. URL: https://opencv.org/ (дата обращения: 22.04.2026).",
        "12) Ultralytics Docs. [Электронный ресурс]. URL: https://docs.ultralytics.com/ (дата обращения: 22.04.2026).",
        "13) ReportLab. [Электронный ресурс]. URL: https://www.reportlab.com/ (дата обращения: 22.04.2026).",
        "14) Локальный репозиторий проекта komarov_thesis. Исходный код и сопроводительная документация.",
    ]
    for item in sources:
        add_paragraph(doc, item, first_line_indent_cm=None)


def add_change_log_page(doc: Document):
    add_heading(doc, "ЛИСТ РЕГИСТРАЦИИ ИЗМЕНЕНИЙ")
    table = doc.add_table(rows=7, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = [
        "Изм.",
        "Номера листов\n(страниц) измененных",
        "Номера листов\n(страниц) замененных",
        "Номера листов\n(страниц) новых",
        "Номера листов\n(страниц) аннулированных",
        "№ документа",
        "Входящий № сопроводительного документа и дата",
        "Подпись,\nдата",
    ]
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, size=11, bold=True)
    for row in table.rows[1:]:
        for cell in row.cells:
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run("")
            set_run_font(run, size=11)


def build_document() -> Document:
    doc = Document()
    configure_document(doc)
    doc.core_properties.author = STUDENT_NAME
    doc.core_properties.title = f"{PROGRAM_NAME_RU}. Текст программы"
    doc.core_properties.subject = "Программная документация"

    add_approval_page(doc)
    add_page_break(doc)
    add_title_page(doc)
    add_page_break(doc)

    add_heading(doc, "СОДЕРЖАНИЕ")
    toc_paragraph = doc.add_paragraph()
    toc_paragraph.paragraph_format.first_line_indent = None
    toc_paragraph.paragraph_format.line_spacing = 1.5
    add_toc(toc_paragraph)

    add_page_break(doc)
    add_introduction(doc)
    add_program_text(doc)
    add_sources(doc)
    add_page_break(doc)
    add_change_log_page(doc)
    return doc


def main():
    doc = build_document()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
