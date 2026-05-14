from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.text.paragraph import Paragraph

from generate_updated_vkr_docx import (
    FONT_BOX,
    FONT_SMALL,
    FONT_TEXT,
    _arrow,
    _box,
    _canvas,
    _title,
)


DOCS_DIR = Path("backend/docs")
WORK_DIR = next(path for path in DOCS_DIR.iterdir() if path.is_dir() and "2026-05-14" in path.name)
SOURCE_DOCX = WORK_DIR / "ВКР_комаров_актуализировано.docx"
OUTPUT_DOCX = WORK_DIR / "ВКР_комаров_актуализировано_v2.docx"
ASSET_DIR = WORK_DIR / "assets_v2"
IMAGE_WIDTH_CM = 15.7


def draw_general_architecture(path: Path) -> None:
    image, draw = _canvas()
    _title(
        draw,
        "Контейнерная архитектура системы",
        "Запуск через docker compose: frontend, backend, PostgreSQL и файловое хранилище",
    )

    _box(draw, (80, 205, 345, 365), "Пользователь\nбраузер", "#eef7ff", "#3c74a6")
    _box(draw, (520, 185, 830, 385), "frontend\nNginx + React build\nпорт 3000", "#f3fbf4", "#4b8f55")
    _box(draw, (1030, 185, 1360, 385), "backend\nFastAPI + Uvicorn\nпорт 8000", "#fff7e8", "#a8782d")
    _box(draw, (1490, 185, 1725, 385), "inline worker\nфоновые задачи", "#fff1f3", "#a04a57")

    _box(draw, (525, 610, 820, 805), "db-seed\nSQLite seed ->\nPostgreSQL", "#f7f2ff", "#7456a6", FONT_TEXT)
    _box(draw, (910, 610, 1220, 805), "postgres_data\nPostgreSQL\nпользователи, проекты,\nоборудование, аудит", "#f7f2ff", "#7456a6", FONT_TEXT)
    _box(draw, (1320, 610, 1685, 805), "app_data -> /data\nuploads, outputs,\ndebug_output,\nstorage_objects", "#f6fafc", "#47738a", FONT_TEXT)
    _box(draw, (130, 610, 410, 805), "storage-init\nкопирует seed-файлы\nодин раз", "#f6fafc", "#47738a", FONT_TEXT)

    _arrow(draw, (345, 285), (520, 285))
    _arrow(draw, (830, 285), (1030, 285))
    _arrow(draw, (1360, 285), (1490, 285))
    _arrow(draw, (1190, 385), (1065, 610))
    _arrow(draw, (665, 610), (910, 705))
    _arrow(draw, (270, 610), (1320, 710))
    _arrow(draw, (1195, 385), (1500, 610))

    draw.text((80, 925), "Внешний доступ: http://localhost:3000; Swagger/отладка: http://localhost:8000.", font=FONT_TEXT, fill="#334155")
    draw.text((80, 965), "Повторные запуски используют Docker volumes и не перезаписывают пользовательские данные.", font=FONT_TEXT, fill="#334155")
    image.save(path)


def draw_backend_components(path: Path) -> None:
    image, draw = _canvas()
    _title(draw, "Компоненты серверной части", "FastAPI, сервисы, вычислительное ядро и инфраструктура хранения")

    layers = [
        ("HTTP/API", 170, "#eef7ff", "#3c74a6", ["routers", "Pydantic schemas", "auth/RBAC", "OpenAPI"]),
        ("Application", 365, "#fff7e8", "#a8782d", ["projects", "uploads", "pipeline", "equipment", "PDF versions", "audit"]),
        ("Processing", 560, "#fffbe8", "#9c8226", ["OpenCV preprocessing", "wall morphology", "OCR dimensions", "YOLO active models", "fire alarm layout"]),
        ("Infrastructure", 755, "#f7f2ff", "#7456a6", ["SQLAlchemy", "PostgreSQL/SQLite", "StorageService", "BackgroundTask", "seed tools"]),
    ]
    for title, y, fill, outline, items in layers:
        _box(draw, (85, y, 355, y + 125), title, fill, outline)
        x = 440
        for item in items:
            _box(draw, (x, y, x + 205, y + 125), item, "#ffffff", outline, FONT_TEXT)
            x += 245
        _arrow(draw, (355, y + 62), (440, y + 62), outline)

    for y in (295, 490, 685):
        _arrow(draw, (930, y), (930, y + 70), "#64748b")

    draw.text((90, 980), "Длительные операции возвращают 202 Accepted и обновляют progress_percent/progress_stage.", font=FONT_TEXT, fill="#334155")
    image.save(path)


def draw_frontend_components(path: Path) -> None:
    image, draw = _canvas()
    _title(draw, "Компоненты клиентской части", "React-приложение, страницы сценариев и общий API-клиент")

    _box(draw, (90, 185, 395, 360), "App shell\nroutes\nprotected pages", "#eef7ff", "#3c74a6")
    _box(draw, (560, 145, 830, 300), "Проекты\nсоздание, фильтры,\nкарточка", "#f3fbf4", "#4b8f55", FONT_TEXT)
    _box(draw, (890, 145, 1160, 300), "Редактор плана\nпроверка результатов\nбез изменений логики", "#f3fbf4", "#4b8f55", FONT_TEXT)
    _box(draw, (1220, 145, 1490, 300), "Оборудование\nкаталог и привязка\nк проекту", "#f3fbf4", "#4b8f55", FONT_TEXT)
    _box(draw, (560, 430, 830, 585), "Фоновые задачи\nстатус, стадия,\nпрогресс", "#fff7e8", "#a8782d", FONT_TEXT)
    _box(draw, (890, 430, 1160, 585), "PDF-версии\nистория и скачивание", "#fff7e8", "#a8782d", FONT_TEXT)
    _box(draw, (1220, 430, 1490, 585), "Дообучение\nзапуски, артефакты,\nактивация модели", "#fff7e8", "#a8782d", FONT_TEXT)
    _box(draw, (720, 730, 1010, 895), "API client\n/api через Nginx proxy\npolling >= 2 сек.", "#f7f2ff", "#7456a6", FONT_TEXT)
    _box(draw, (1090, 730, 1380, 895), "Админ-раздел\nпользователи и аудит\nrole=developer", "#f7f2ff", "#7456a6", FONT_TEXT)

    for start in [(395, 272), (695, 300), (1025, 300), (1355, 300), (695, 585), (1025, 585), (1355, 585), (1090, 812)]:
        _arrow(draw, start, (720, 812), "#64748b")

    image.save(path)


def draw_database_schema(path: Path) -> None:
    image, draw = _canvas()
    _title(draw, "Логическая схема хранения данных", "Реляционная БД хранит метаданные; файлы вынесены в DATA_DIR/app_data")

    tables = [
        ("users", "id\nusername\npassword_hash\nrole", 70, 160),
        ("projects", "id\nauthor_id\nobject_name\nnumber\nstatus\nis_deleted", 420, 160),
        ("floor_plans", "id\nproject_id\nsource_path\nwidth\nheight\nscale", 785, 160),
        ("equipment_items", "id\nname\ncategory\nprice\nimage_path", 1150, 160),
        ("background_tasks", "id\nproject_id\nkind\nstatus\nprogress", 70, 520),
        ("audit_events", "id\nuser_id\naction\nentity\ntimestamp", 420, 520),
        ("project_pdf_versions", "id\nproject_id\nfile_path\nauthor_id\nparams", 785, 520),
        ("recognition_training_runs", "id\ntarget\nstatus\nmetrics\nartifact_paths", 1150, 520),
    ]
    for name, fields, x, y in tables:
        _box(draw, (x, y, x + 265, y + 245), f"{name}\n{fields}", "#ffffff", "#475569", FONT_SMALL, radius=10)

    for start, end in [
        ((335, 280), (420, 280)),
        ((685, 280), (785, 280)),
        ((552, 405), (552, 520)),
        ((552, 280), (70, 640)),
        ((552, 280), (785, 640)),
        ((552, 280), (1150, 640)),
    ]:
        _arrow(draw, start, end, "#64748b", width=3)

    _box(draw, (1435, 270, 1715, 475), "postgres_data\nPostgreSQL volume\nили SQLite файл\nпри локальном запуске", "#f7f2ff", "#7456a6", FONT_TEXT)
    _box(draw, (1435, 600, 1715, 820), "app_data (/data)\nuploads\noutputs\ndebug_output\nstorage_objects", "#f6fafc", "#47738a", FONT_TEXT)
    _arrow(draw, (1305, 280), (1435, 360), "#64748b")
    _arrow(draw, (1050, 700), (1435, 710), "#64748b")
    image.save(path)


def draw_class_diagram(path: Path) -> None:
    image, draw = _canvas()
    _title(draw, "Диаграмма классов и ключевых сущностей", "Предметные модели, служебные процессы и артефакты проекта")

    classes = [
        ("Project", "author\nstatus\nsoft delete", 80, 175),
        ("FloorPlan", "source_path\nscale\nimage size", 390, 175),
        ("Wall", "x1,y1,x2,y2\nthickness\nlength", 700, 175),
        ("Opening", "wall_id\ntype\nposition", 1010, 175),
        ("Room", "contour\narea\nlabel", 1320, 175),
        ("EquipmentItem", "category\nprice\nasset paths", 80, 520),
        ("ProjectPdfVersion", "file_path\nauthor\nparams", 390, 520),
        ("BackgroundTask", "status\nprogress\nretry", 700, 520),
        ("AuditEvent", "user\naction\nentity", 1010, 520),
        ("RecognitionTrainingRun", "target\nmetrics\nartifacts", 1320, 520),
    ]
    for title, fields, x, y in classes:
        _box(draw, (x, y, x + 245, y + 210), f"{title}\n{fields}", "#ffffff", "#475569", FONT_TEXT, radius=10)

    for start, end in [
        ((325, 280), (390, 280)),
        ((635, 280), (700, 280)),
        ((945, 280), (1010, 280)),
        ((1255, 280), (1320, 280)),
        ((202, 385), (202, 520)),
        ((512, 385), (512, 520)),
        ((822, 385), (822, 520)),
        ((1132, 385), (1132, 520)),
        ((1442, 385), (1442, 520)),
    ]:
        _arrow(draw, start, end, "#64748b", width=3)
    image.save(path)


def generate_figures() -> dict[int, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    figures = {
        1: ASSET_DIR / "figure_1_general_architecture_v2.png",
        2: ASSET_DIR / "figure_2_backend_components_v2.png",
        3: ASSET_DIR / "figure_3_frontend_components_v2.png",
        4: ASSET_DIR / "figure_4_database_schema_v2.png",
        5: ASSET_DIR / "figure_5_class_diagram_v2.png",
    }
    draw_general_architecture(figures[1])
    draw_backend_components(figures[2])
    draw_frontend_components(figures[3])
    draw_database_schema(figures[4])
    draw_class_diagram(figures[5])
    return figures


def clear_runs(paragraph: Paragraph) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_text(paragraph: Paragraph, text: str) -> None:
    clear_runs(paragraph)
    paragraph.add_run(text)


def insert_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    if style:
        new_paragraph.style = style
    new_paragraph.add_run(text)
    return new_paragraph


def find_paragraph(doc: Document, text: str, *, start: int = 0) -> tuple[int, Paragraph]:
    for index, paragraph in enumerate(doc.paragraphs[start:], start=start):
        if paragraph.text.strip() == text:
            return index, paragraph
    raise ValueError(f"Paragraph not found: {text}")


def find_startswith(doc: Document, prefix: str, *, start: int = 0) -> tuple[int, Paragraph]:
    for index, paragraph in enumerate(doc.paragraphs[start:], start=start):
        if paragraph.text.strip().startswith(prefix):
            return index, paragraph
    raise ValueError(f"Paragraph not found by prefix: {prefix}")


def replace_figure(paragraph: Paragraph, image_path: Path) -> None:
    clear_runs(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(IMAGE_WIDTH_CM))


def replace_figures(doc: Document, figures: dict[int, Path]) -> None:
    for number, image_path in figures.items():
        caption_index, _ = find_startswith(doc, f"Рисунок {number}")
        replace_figure(doc.paragraphs[caption_index - 1], image_path)


def update_storage_sections(doc: Document) -> None:
    storage_idx, _ = find_paragraph(doc, "Хранение данных")
    set_text(
        doc.paragraphs[storage_idx + 1],
        "Для хранения данных используется реляционная модель. При локальном запуске базовой конфигурацией остается SQLite: файл floor_plans.db располагается в каталоге данных и позволяет быстро развернуть демонстрационную версию без отдельного сервера БД. Для контейнерного и более приближенного к промышленной эксплуатации запуска используется PostgreSQL, подключаемый через переменную окружения DATABASE_URL. Прикладной код работает через SQLAlchemy, поэтому выбранная СУБД не меняет публичные сценарии системы.",
    )
    set_text(
        doc.paragraphs[storage_idx + 2],
        "Рабочие файлы отделены от исходного кода и задаются параметром DATA_DIR. Внутри каталога данных используются подкаталоги uploads для исходных планов, outputs для результатов обработки и PDF-документов, debug_output для диагностических материалов и storage_objects для объектного хранилища приложения. В базе данных сохраняются метаданные и нормализованные относительные пути, а не содержимое файлов.",
    )
    set_text(
        doc.paragraphs[storage_idx + 3],
        "В Docker-конфигурации реляционные данные хранятся в named volume postgres_data, а файловое хранилище подключается как named volume app_data и монтируется в backend-контейнер как /data. Первичное наполнение выполняется одноразовыми сервисами: db-seed переносит снимок SQLite в PostgreSQL, а storage-init копирует seed-файлы и создает marker-файл, чтобы повторные запуски не перезаписывали пользовательские добавления.",
    )

    arch_idx, _ = find_paragraph(doc, "Архитектура хранения данных")
    set_text(
        doc.paragraphs[arch_idx + 1],
        "База данных содержит сведения о пользователях, проектах, планах этажей, оборудовании, фоновых задачах, версиях PDF, событиях аудита и запусках дообучения распознавания. Сущность проекта является центральной: с ней связаны загруженные планы, результаты обработки, сформированные документы, выбранное оборудование, служебные задачи и записи аудита.",
    )
    set_text(
        doc.paragraphs[arch_idx + 2],
        "Для предметных данных используется SQLAlchemy-модель, совместимая с SQLite и PostgreSQL. Локальный запуск сохраняет простоту SQLite, а Docker-запуск использует PostgreSQL в volume postgres_data. Это позволяет демонстрировать систему одной командой docker compose up --build и при этом не терять данные после остановки контейнеров.",
    )
    set_text(
        doc.paragraphs[arch_idx + 3],
        "Файлы планов, PDF-версии, изображения оборудования, логи и артефакты дообучения физически располагаются вне БД. Для них используется файловое хранилище DATA_DIR, в Docker — volume app_data, смонтированный как /data. В таблицах остаются относительные пути и параметры формирования, поэтому перенос проекта между локальным и контейнерным запуском не требует сохранения абсолютных путей Windows.",
    )

    notes_idx, _ = find_paragraph(doc, "2.3.2. Примечания к хранению файлов и оптимизации")
    set_text(
        doc.paragraphs[notes_idx + 1],
        "Файлы планов и результаты обработки сохраняются отдельно от исходного кода. В базе данных хранится только нормализованный относительный путь: это предотвращает ошибки при переносе между Windows-путями, локальным запуском из backend и контейнерной директорией /data. Публичные ссылки строятся через backend-механизм ассетов, который не раскрывает абсолютные пути файловой системы.",
    )
    set_text(
        doc.paragraphs[notes_idx + 2],
        "Для изображений применяется серверная проверка расширения, MIME-типа, размера файла и минимального разрешения 200x200 пикселей. Проверка на стороне frontend используется для ранней обратной связи пользователю, однако окончательное решение принимает backend. Это важно, потому что контейнерный frontend раздается Nginx как статическое приложение и не может считаться доверенной стороной в вопросах валидации файлов.",
    )
    insert_after(
        doc.paragraphs[notes_idx + 2],
        "Оптимизация хранения основана на разделении больших бинарных артефактов и реляционных данных. PostgreSQL хранит связи, статусы, параметры, версии и аудит, а app_data хранит сами файлы. При первом запуске Docker seed копируется только в пустой volume; дальнейшие изменения остаются индивидуальными для конкретного развертывания и сохраняются после docker compose down.",
        style="t3",
    )


def update_wall_algorithm(doc: Document) -> None:
    heading_idx, _ = find_paragraph(doc, "Алгоритм распознавания стен")
    texts = [
        "Распознавание стен начинается с подготовки изображения средствами OpenCV. На этапе rectify_document изображение переводится в оттенки серого, сглаживается GaussianBlur, после чего границы ищутся алгоритмом Canny. По найденным границам выделяются внешние контуры; для крупного контура выполняется аппроксимация многоугольника approxPolyDP, а если явный четырехугольник не найден, используется minAreaRect. Далее рассчитывается матрица перспективного преобразования getPerspectiveTransform и план приводится к прямоугольному виду через warpPerspective.",
        "При необходимости применяется deskew: по Canny-границам вызывается HoughLinesP, из длинных почти горизонтальных линий вычисляется медианный угол наклона, после чего изображение поворачивается warpAffine. Затем функция binarize переводит план в рабочую бинарную маску: используется grayscale, medianBlur, оценка фона через GaussianBlur, нормализация яркости операцией divide и adaptiveThreshold с гауссовым окном. После этого выполняются morphology close/open и удаление мелких connected components.",
        "Базовый алгоритм detect_walls работает с бинарной маской. Сначала ensure_foreground_white приводит изображение к виду «стены белым по черному фону». Затем preprocess_wall_mask усиливает области стен: применяется морфологическое закрытие прямоугольным ядром и dilate, чтобы тонкие и разорванные линии стали связными областями.",
        "После подготовки маска разделяется на горизонтальные и вертикальные кандидаты. Для этого extract_horizontal_vertical_masks использует morphology open с длинным горизонтальным прямоугольным ядром и отдельным вертикальным ядром. По каждой маске вызывается connectedComponentsWithStats: компоненты фильтруются по площади, минимальной длине и допустимой толщине. Толщина уточняется через distanceTransform, который оценивает максимальный радиус внутри компоненты.",
        "Каждая подтвержденная компонента преобразуется в объект Wall с осевой линией, углом 0 или 90 градусов, толщиной и confidence. Фрагменты одной стены объединяются функцией merge_walls, если они коллинеарны, имеют близкое смещение и небольшой разрыв между концами. Если для класса walls активирована дообученная YOLO-модель, сервис сначала использует предсказания Ultralytics: mask-полигоны или bounding boxes преобразуются в стены. Если активной модели нет, используется описанный OpenCV-морфологический алгоритм.",
    ]
    for offset, text in enumerate(texts[:3], start=1):
        set_text(doc.paragraphs[heading_idx + offset], text)
    anchor = doc.paragraphs[heading_idx + 3]
    for text in reversed(texts[3:]):
        insert_after(anchor, text, style="t3")


def update_use_cases(doc: Document) -> None:
    _, conclusions = find_paragraph(doc, "Выводы", start=find_paragraph(doc, "Пользовательские сценарии (UseCases)")[0])
    items = [
        ("2.4.7. UC-7: Авторизация и разграничение ролей", "Пользователь открывает систему и проходит авторизацию по имени пользователя и паролю. Backend проверяет PBKDF2-хэш пароля и выдает доступ в соответствии с ролью. Роль engineer предназначена для рабочих инженерных сценариев, а роль developer используется как текущий эквивалент администратора и открывает управление пользователями, аудитом и расширенными административными операциями."),
        ("2.4.8. UC-8: Ведение каталога оборудования", "Пользователь с правами developer ведет каталог оборудования: добавляет позиции, категории, изображения, схемы подключения и стоимость. Инженер использует каталог при наполнении проекта, выбирает устройства для плана и получает согласованную спецификацию оборудования в проектной документации."),
        ("2.4.9. UC-9: Дообучение распознавания", "Администратор просматривает накопленные примеры корректировок, исключает неподходящие данные и запускает дообучение модели распознавания. Система создает фоновую задачу, сохраняет логи, метрики и артефакты обучения. После проверки результата модель может быть активирована для последующих запусков распознавания."),
        ("2.4.10. UC-10: Аудит действий и экспорт журнала", "Пользователь с ролью developer открывает страницу аудита, фильтрует события по пользователю, действию, сущности и периоду. При необходимости журнал выгружается в CSV для табличного анализа или в JSON для машинной обработки. Это позволяет восстановить последовательность действий при проверке проекта."),
        ("2.4.11. UC-11: Контроль фоновых задач", "Пользователь отслеживает длительные операции через список задач и индикаторы прогресса. Для каждой задачи отображаются статус, процент выполнения, текущая стадия, сообщение об ошибке и связь с проектом. Если задача завершилась ошибкой и тип операции допускает повторный запуск, пользователь инициирует retry без удаления истории предыдущей попытки."),
        ("2.4.12. UC-12: Контейнерное развертывание со стандартными данными", "Оператор развертывает систему командой docker compose up --build. При первом запуске PostgreSQL получает seed-снимок текущих пользователей, проектов, оборудования и задач, а файловое хранилище app_data получает связанные uploads, outputs, debug_output и storage_objects. После этого каждое развертывание хранит новые данные в собственных Docker volumes."),
    ]
    for heading, body in items:
        heading_paragraph = conclusions.insert_paragraph_before(heading)
        heading_paragraph.style = "p3"
        body_paragraph = conclusions.insert_paragraph_before(body)
        body_paragraph.style = "t3"


def validate(doc: Document) -> None:
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    bad_tokens = ["Рџ", "Рљ", "вЂ"]
    found = [token for token in bad_tokens if token in text]
    if found:
        raise RuntimeError(f"Possible mojibake tokens found: {found}")
    if len(doc.inline_shapes) != 5:
        raise RuntimeError(f"Expected 5 inline figures, found {len(doc.inline_shapes)}")


def main() -> None:
    figures = generate_figures()
    doc = Document(SOURCE_DOCX)
    replace_figures(doc, figures)
    update_storage_sections(doc)
    update_wall_algorithm(doc)
    update_use_cases(doc)
    validate(doc)
    doc.save(OUTPUT_DOCX)
    print(f"Saved {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
